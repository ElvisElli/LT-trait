"""
GC2 raw weight data → GC2_clean.csv
GC Experiment 2: PI398245 (Q1), PI567488B (Q2), PI377574 (Q3), PI424397 (Q4)
6 VPD treatments, 4 replications, 2 days (10/30–10/31)
Weight loss between start and end of each treatment = transpiration (ml)
"""

import openpyxl
import pandas as pd
import os

XL = "experimental_data/raw_data/GC2/GC_experiment_2.xlsx"
wb = openpyxl.load_workbook(XL, data_only=True)

# ── 1. Pot → genotype / quadrant mapping ────────────────────────────────────
pot_info = {
    1:  ("PI398245",  "Q1", "Plastic, High WUE"),
    2:  ("PI567488B", "Q2", "Plastic, Low WUE"),
    3:  ("PI377574",  "Q3", "Nonplastic, Low WUE"),
    4:  ("PI424397",  "Q4", "Nonplastic, High WUE"),
    5:  ("PI567488B", "Q2", "Plastic, Low WUE"),
    6:  ("PI377574",  "Q3", "Nonplastic, Low WUE"),
    7:  ("PI398245",  "Q1", "Plastic, High WUE"),
    8:  ("PI424397",  "Q4", "Nonplastic, High WUE"),
    9:  ("PI424397",  "Q4", "Nonplastic, High WUE"),
    10: ("PI567488B", "Q2", "Plastic, Low WUE"),
    11: ("PI377574",  "Q3", "Nonplastic, Low WUE"),
    12: ("PI398245",  "Q1", "Plastic, High WUE"),
    13: ("PI377574",  "Q3", "Nonplastic, Low WUE"),
    14: ("PI398245",  "Q1", "Plastic, High WUE"),
    15: ("PI424397",  "Q4", "Nonplastic, High WUE"),
    16: ("PI567488B", "Q2", "Plastic, Low WUE"),
}

# Replication number within each genotype (ordered by pot number)
geno_pots = {}
for pot, (geno, q, wue) in pot_info.items():
    geno_pots.setdefault(geno, []).append(pot)
for geno in geno_pots:
    geno_pots[geno].sort()

pot_rep = {}
for geno, pots in geno_pots.items():
    for rep_idx, pot in enumerate(pots, start=1):
        pot_rep[pot] = rep_idx

# ── 2. VPD target values from Evaluation targets sheet ──────────────────────
vpd_targets = {1: 0.951, 2: 1.522, 3: 2.092, 4: 2.663, 5: 3.233, 6: 3.804}

# Actual average VPD from Tracking VPD (col 9 = Day1, col 20 = Day2; rows 2–7 = T1–T6)
ws_vpd = wb["Tracking VPD"]
vpd_rows = list(ws_vpd.iter_rows(values_only=True))
# rows index 1–6 → treatments T1–T6
vpd_actual = {}
for t_idx in range(6):
    row = vpd_rows[t_idx + 1]  # 0-indexed, skip header
    vpd_actual[(t_idx + 1, 1)] = row[8]   # Day1 average
    vpd_actual[(t_idx + 1, 2)] = row[19]  # Day2 average

# Temperature and RH: target values (32°C, decreasing RH)
# From Evaluation targets rows 12–16 (0-indexed 11–15)
ws_ev = wb["Evaluation targets"]
ev_rows = list(ws_ev.iter_rows(values_only=True))
# Row 12 (idx 11): RH Day1: col3=pre, col4=T1..col9=T6
# Row 13 (idx 12): Temp Day1
# Row 17 (idx 16): treatment labels T1..T6 (cols 3–8)
rh_day1  = [ev_rows[11][c] for c in range(3, 9)]   # T1..T6
rh_day2  = [ev_rows[11][c] for c in range(12, 17)]  # T1..T5 (T6 missing→use T5)
rh_day2.append(rh_day2[-1])  # T6 ≈ T5 (chamber was below target)

temp_day1 = [ev_rows[12][c] for c in range(3, 9)]
temp_day2 = [ev_rows[12][c] for c in range(12, 17)]
temp_day2.append(temp_day2[-1])

# ── 3. Extract weights from Tracking weights ─────────────────────────────────
ws_tw = wb["Tracking weights"]
tw_rows = list(ws_tw.iter_rows(values_only=True))  # 0-indexed; row 0 = header

# Block layout: 16 rows per block (pots 1–16), alternating Start / End per treatment
# T1 Start: rows 1–16, T1 End: 17–32, T2 Start: 33–48, T2 End: 49–64, ...
def extract_weights(day_col_total, day_col_weight):
    """Return dict {(pot, treatment): (start_w, end_w)}."""
    treatment_row_starts = {
        1: (1,  17),   # (start_block_idx, end_block_idx) in tw_rows
        2: (33, 49),
        3: (65, 81),
        4: (97, 113),
        5: (129, 145),
        6: (161, 177),
    }
    result = {}
    for t, (s_idx, e_idx) in treatment_row_starts.items():
        for offset in range(16):
            pot = tw_rows[s_idx + offset][2]   # pot number (col C = index 2)
            w_start = tw_rows[s_idx + offset][day_col_weight]
            w_end   = tw_rows[e_idx + offset][day_col_weight]
            if pot is not None and w_start is not None and w_end is not None:
                result[(int(pot), t)] = (float(w_start), float(w_end))
    return result

weights_d1 = extract_weights(3, 4)   # col D=tare(3), col E=weight(4)
weights_d2 = extract_weights(16, 17) # col Q=tare(16), col R=weight(17)

# ── 4. Compute transpiration and relative transpiration ──────────────────────
def build_df(weights_by_day, day_num, rh_list, temp_list):
    records = []
    # First pass: compute absolute transpiration
    transp = {}
    for (pot, t), (ws, we) in weights_by_day.items():
        transp[(pot, t)] = ws - we  # ml (weight loss = water loss)

    # Normalise per pot: RelTransp = T_i / max(T_all_treatments)
    pots = sorted(set(p for p, _ in transp.keys()))
    for pot in pots:
        t_vals = {t: transp.get((pot, t), None) for t in range(1, 7)}
        valid = [v for v in t_vals.values() if v is not None]
        if not valid:
            continue
        max_t = max(valid)

        geno, quad, wue_class = pot_info[pot]
        rep = pot_rep[pot]
        plasticity = "Plastic" if "Plastic" in wue_class else "Nonplastic"
        wue_simple = "High WUE" if "High" in wue_class else "Low WUE"

        for t in range(1, 7):
            if (pot, t) not in weights_by_day:
                continue
            abs_t = transp[(pot, t)]
            # Negative transpiration is physically impossible — mark as NA
            if abs_t <= 0:
                abs_t = None
            valid2 = [v for v in t_vals.values() if v is not None and v > 0]
            max_t2 = max(valid2) if valid2 else None
            rel_t = (abs_t / max_t2) if (abs_t is not None and max_t2 and max_t2 > 0) else None
            records.append({
                "Genotype":             geno,
                "Quadrant":             quad,
                "Expected_WUE":         wue_class,
                "Replication":          rep,
                "Pot":                  pot,
                "Day":                  day_num,
                "Treatment":            t,
                "Temperature":          temp_list[t - 1],
                "RH":                   rh_list[t - 1],
                "VPD":                  vpd_actual.get((t, day_num)),
                "RelTranspiration":     rel_t,
                "Transpiration_ml":     abs_t,  # None = anomalous (irrigation artifact)
                "TranspLeafArea":       None,
                "DryMatter_leaf":       None,
                "DryMatter_stem":       None,
                "DryMatter_total":      None,
                "LeafArea_squarecm":    None,
                "StomatalConductance":  None,
                "E_apparent":           None,
                "T_leaf_degreecelcius": None,
                "Experiment":           "GC2",
                "DryMatter_rep":        None,
                "VPD_target":           vpd_targets[t],
                "Plasticity":           plasticity,
                "WUE_class":            wue_simple,
            })
    return pd.DataFrame(records)

df_d1 = build_df(weights_d1, 1, rh_day1, temp_day1)
df_d2 = build_df(weights_d2, 2, rh_day2, temp_day2)
gc2 = pd.concat([df_d1, df_d2], ignore_index=True)

# Sort for readability
gc2.sort_values(["Genotype", "Replication", "Day", "Treatment"], inplace=True)
gc2.reset_index(drop=True, inplace=True)

out_dir = "experimental_data/cleaned_data"
os.makedirs(out_dir, exist_ok=True)
gc2.to_csv(f"{out_dir}/GC2_clean.csv", index=False)
print(f"GC2_clean.csv written: {len(gc2)} rows")
print(gc2.groupby(["Genotype", "Day"])["RelTranspiration"].describe())
print("\nSample rows:")
print(gc2[gc2["Genotype"] == "PI398245"].head(12).to_string())
